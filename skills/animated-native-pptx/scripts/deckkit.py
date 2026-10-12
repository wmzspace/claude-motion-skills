"""deckkit: native-animated, editable PPTX on a 1920x1080 px grid (1 px = 0.5 pt)."""
import subprocess, zipfile, os
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from lxml import etree

PX = 6350
L, C, R = PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT
MID = MSO_ANCHOR.MIDDLE
NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
T = {}          # theme: font, on_bg, on_bg_soft, accent, accent_ink, card, ink, mute, primary, dot
A = []          # animations queued for the current slide
AUD = {}        # slide id -> (shape id, duration ms) for auto-playing audio
S = {'prs': None, 'n': 0, 'total': 0, 'tag': ''}
e = lambda v: Emu(int(v * PX))
rgb = lambda h: RGBColor.from_string(h)


def init(theme, total, tag=''):
    T.update(theme); S.update(total=total, tag=tag, n=0)
    S['prs'] = Presentation(); S['prs'].slide_width = e(1920); S['prs'].slide_height = e(1080)
    return S['prs']


def add(kind, shp, delay, dur=600, **kw):
    """Queue an animation (ms). kind: fade rise left right zoom wipe pulse path."""
    A.append((delay, kind, shp.shape_id, dur, kw)); return shp


def box(c, x, y, w, h, fill, kind=MSO_SHAPE.ROUNDED_RECTANGLE, rad=40, alpha=None, name=None):
    s = c.add_shape(kind, e(x), e(y), e(w), e(h))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE: s.adjustments[0] = min(.5, rad / min(w, h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill); s.line.fill.background(); s.shadow.inherit = False
    if alpha is not None:
        etree.SubElement(s._element.spPr.find('.//' + qn('a:srgbClr')), qn('a:alpha')).set('val', str(int(alpha * 100000)))
    if name: s.name = name
    return s


def style(tf, t, pt, color, bold=False, align=L, bullets=False, spacing=None):
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(t.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing: p.line_spacing = spacing
        r = p.add_run(); r.text = line
        r.font.size = Pt(pt); r.font.bold = bold; r.font.color.rgb = rgb(color); r.font.name = T['font']
        etree.SubElement(r._r.get_or_add_rPr(), qn('a:ea')).set('typeface', T['font'])
        if bullets:
            pPr = p._p.get_or_add_pPr(); pPr.set('marL', '342900'); pPr.set('indent', '-342900')
            etree.SubElement(pPr, qn('a:buFont')).set('typeface', 'Arial')
            etree.SubElement(pPr, qn('a:buChar')).set('char', '•')


def text(c, x, y, w, h, t, pt, color=None, bold=False, align=L, anchor=MSO_ANCHOR.TOP, **kw):
    s = c.add_textbox(e(x), e(y), e(w), e(h)); s.text_frame.vertical_anchor = anchor
    style(s.text_frame, t, pt, color or T['on_bg'], bold, align, **kw); return s


def pill(c, x, y, t, pt=13, fill=None, color=None, w=None, h=46, center=False):
    w = w or len(t) * pt * 2 + 56
    if center: x = x - w / 2
    s = box(c, x, y, w, h, fill or T['accent'], rad=h / 2); s.text_frame.vertical_anchor = MID
    style(s.text_frame, t, pt, color or T['accent_ink'], True, C); return s


def pic(c, path, x, y, w, h=None):
    return c.add_picture(path, e(x), e(y), e(w), e(h or w))


def photo(c, path, x, y, w, h, contain=False, rad=6000, border=True):
    """Photo in a rounded frame. Default crops to fill (cover); contain=True fits inside."""
    iw, ih = Image.open(path).size
    if contain:
        s = min(w / iw, h / ih); pw, ph = iw * s, ih * s
        p = c.add_picture(path, e(x + (w - pw) / 2), e(y + (h - ph) / 2), e(pw), e(ph))
    else:
        p = c.add_picture(path, e(x), e(y), e(w), e(h)); r = (iw / ih) / (w / h)
        if r > 1: p.crop_left = p.crop_right = (1 - 1 / r) / 2
        else: p.crop_top = p.crop_bottom = (1 - r) / 2
    g = p._element.spPr.find(qn('a:prstGeom')); g.set('prst', 'roundRect')
    av = g.find(qn('a:avLst'))
    if av is None: av = etree.SubElement(g, qn('a:avLst'))
    gd = etree.SubElement(av, qn('a:gd')); gd.set('name', 'adj'); gd.set('fmla', f'val {rad}')
    if border: p.line.color.rgb = rgb('FFFFFF'); p.line.width = Pt(3)
    return p


def badge(c, icon, x, y, s=150, bg='FFFFFF', ic='primary'):
    """Circle with a line icon; needs assets/<icon>_<ic>.png from assets.js."""
    box(c, x, y, s, s, bg, MSO_SHAPE.OVAL); k = s * .58
    pic(c, f'assets/{icon}_{ic}.png', x + (s - k) / 2, y + (s - k) / 2, k)


def group(sl, name=None):
    g = sl.shapes.add_group_shape()
    if name: g.name = name
    return g


def card(sl, x, y, w, h, name=None):
    """A group holding a translucent rounded card; add content to g.shapes, animate g."""
    g = group(sl, name); box(g.shapes, x, y, w, h, T['card'], alpha=.94); return g


def new(bg, label):
    """New slide: background image assets/<bg>.jpg, top-right tag, page label, progress bar."""
    A.clear(); idx = S['n']; S['n'] += 1; N = S['total']
    sl = S['prs'].slides.add_slide(S['prs'].slide_layouts[6])
    _, rid = sl.part.get_or_add_image_part(f'assets/{bg}.jpg')
    sl._element.find(qn('p:cSld')).insert(0, etree.fromstring(
        f'<p:bg {NS} xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<p:bgPr><a:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>'))
    if S['tag']: text(sl.shapes, 1092, 38, 700, 34, S['tag'], 12, T['on_bg_soft'], True, R)
    text(sl.shapes, 128, 984, 900, 32, f'{idx + 1:02d} / {N:02d}　{label}', 12, T['on_bg'], True)
    sw = (1664 - (N - 1) * 10) / N
    for i in range(N):
        box(sl.shapes, 128 + i * (sw + 10), 1024, sw, 8, T['accent'] if i <= idx else T['on_bg'], rad=4, alpha=None if i <= idx else .28)
    return sl


def header(sl, eyebrow, title, sub='', center=False):
    """Eyebrow pill + title + subtitle. Body area is y 326-950, x 128-1792."""
    if center:
        add('rise', pill(sl.shapes, 960, 90, eyebrow, center=True), 0)
        add('rise', text(sl.shapes, 128, 150, 1664, 100, title, 38, bold=True, align=C), 150); return
    add('zoom', box(sl.shapes, 102, 132, 116, 116, T['dot'], MSO_SHAPE.OVAL), 0, 500)
    add('rise', pill(sl.shapes, 128, 84, eyebrow), 0)
    add('rise', text(sl.shapes, 128, 146, 1600, 100, title, 38, bold=True), 150)
    if sub: add('rise', text(sl.shapes, 128, 252, 1600, 50, sub, 16, T['on_bg_soft']), 350)


def curve(sl, pts, color=None, width=2.5, dash=True):
    """Polyline through pts [(x, y)...] in px; returns (shape, motion-path string for a dot starting at pts[0])."""
    fb = sl.shapes.build_freeform(e(pts[0][0]), e(pts[0][1])); fb.add_line_segments([(e(x), e(y)) for x, y in pts[1:]], close=False)
    ln = fb.convert_to_shape(); ln.fill.background(); ln.line.color.rgb = rgb(color or T['on_bg']); ln.line.width = Pt(width); ln.shadow.inherit = False
    if dash: ln.line.dash_style = MSO_LINE.DASH
    return ln, 'M 0 0 ' + ' '.join(f'L {(x - pts[0][0]) / 1920:.4f} {(y - pts[0][1]) / 1080:.4f}' for x, y in pts[1:])


def audio(sl, path, x=1800, y=950, size=64):
    """Embed an mp3/wav that auto-plays when the slide is entered. Call before finish(); save with save()."""
    icon = 'assets/_speaker.png'
    if not os.path.exists(icon):
        os.makedirs('assets', exist_ok=True); im = Image.new('RGBA', (200, 200), (0, 0, 0, 0)); d = ImageDraw.Draw(im)
        d.ellipse((0, 0, 199, 199), fill=(20, 30, 50, 150)); d.polygon([(50, 80), (80, 80), (115, 50), (115, 150), (80, 120), (50, 120)], fill=(255, 255, 255, 255))
        d.arc((105, 65, 150, 135), -60, 60, fill=(255, 255, 255, 255), width=9); d.arc((95, 45, 175, 155), -55, 55, fill=(255, 255, 255, 255), width=9); im.save(icon)
    try: dur = int(float(subprocess.run(['ffprobe', '-v', 'error', '-show_entries', 'format=duration', '-of', 'csv=p=0', path], capture_output=True, text=True).stdout) * 1000)
    except Exception: dur = 60000
    mime = 'audio/wav' if path.lower().endswith('.wav') else 'audio/mpeg'
    mv = sl.shapes.add_movie(path, e(x), e(y), e(size), e(size), poster_frame_image=icon, mime_type=mime); mv.name = 'audio'
    for t in sl._element.findall(qn('p:timing')): sl._element.remove(t)     # python-pptx adds its own; finish() writes ours
    mv._element.find('.//' + qn('a:videoFile')).tag = qn('a:audioFile')
    AUD[id(sl)] = (mv.shape_id, dur); return mv


def save(path):
    """Save the deck, fixing the relationship type of embedded audio (python-pptx writes it as video)."""
    tmp = path + '.tmp'; S['prs'].save(tmp)
    zi = zipfile.ZipFile(tmp); zo = zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED)
    for it in zi.infolist():
        b = zi.read(it.filename)
        if it.filename.startswith('ppt/slides/_rels/') and (b'.mp3' in b or b'.wav' in b):
            b = b.replace(b'officeDocument/2006/relationships/video', b'officeDocument/2006/relationships/audio')
        zo.writestr(it, b, zipfile.ZIP_STORED if it.filename.lower().endswith(('.mp3', '.wav', '.jpg', '.jpeg', '.png', '.mp4')) else zipfile.ZIP_DEFLATED)
    zo.close(); zi.close(); os.remove(tmp)


def finish(sl, trans='<p:fade/>'):
    """Write the slide transition and all queued animations (auto-play on slide entry)."""
    n = [4]
    def nid(): n[0] += 1; return n[0]
    tgt = lambda s: f'<p:tgtEl><p:spTgt spid="{s}"/></p:tgtEl>'
    def vis(s): return (f'<p:set><p:cBhvr><p:cTn id="{nid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt(s)}'
                        f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>')
    def fx(s, f, d): return f'<p:animEffect transition="in" filter="{f}"><p:cBhvr><p:cTn id="{nid()}" dur="{d}"/>{tgt(s)}</p:cBhvr></p:animEffect>'
    def an(s, attr, frm, d):
        v0 = '<p:fltVal val="0"/>' if frm == '0' else f'<p:strVal val="{frm}"/>'
        return (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn id="{nid()}" dur="{d}" decel="100000" fill="hold"/>{tgt(s)}'
                f'<p:attrNameLst><p:attrName>{attr}</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav tm="0"><p:val>{v0}</p:val></p:tav>'
                f'<p:tav tm="100000"><p:val><p:strVal val="#{attr}"/></p:val></p:tav></p:tavLst></p:anim>')
    pars = []
    for i, (delay, kind, s, d, kw) in enumerate(sorted(A, key=lambda a: a[0])):
        cls, pid, sub = 'entr', 10, 0
        if kind == 'fade': body = vis(s) + fx(s, 'fade', d)
        elif kind == 'rise': pid = 42; body = vis(s) + fx(s, 'fade', d) + an(s, 'ppt_x', '#ppt_x', d) + an(s, 'ppt_y', '#ppt_y+.08', d)
        elif kind == 'left': pid, sub = 2, 8; body = vis(s) + fx(s, 'fade', d) + an(s, 'ppt_x', '#ppt_x-.1', d) + an(s, 'ppt_y', '#ppt_y', d)
        elif kind == 'right': pid, sub = 2, 2; body = vis(s) + fx(s, 'fade', d) + an(s, 'ppt_x', '#ppt_x+.1', d) + an(s, 'ppt_y', '#ppt_y', d)
        elif kind == 'zoom': pid, sub = 53, 16; body = vis(s) + an(s, 'ppt_w', '0', d) + an(s, 'ppt_h', '0', d) + fx(s, 'fade', d)
        elif kind == 'wipe':
            dr = kw.get('dir', 'left'); pid, sub = 22, {'left': 8, 'up': 1, 'right': 2, 'down': 4}[dr]; body = vis(s) + fx(s, f'wipe({dr})', d)
        elif kind == 'pulse':
            cls, pid = 'emph', 6
            body = f'<p:animScale><p:cBhvr><p:cTn id="{nid()}" dur="{d}" autoRev="1" fill="hold"/>{tgt(s)}</p:cBhvr><p:by x="108000" y="108000"/></p:animScale>'
        elif kind == 'path':
            cls, pid = 'path', 0
            body = (f'<p:animMotion origin="layout" path="{kw["path"]}" pathEditMode="relative" ptsTypes=""><p:cBhvr><p:cTn id="{nid()}" dur="{d}" fill="hold"/>{tgt(s)}'
                    f'<p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr></p:animMotion>')
        else: raise ValueError(kind)
        pars.append(f'<p:par><p:cTn id="{nid()}" presetID="{pid}" presetClass="{cls}" presetSubtype="{sub}" fill="hold" nodeType="{"afterEffect" if i == 0 else "withEffect"}">'
                    f'<p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>{body}</p:childTnLst></p:cTn></p:par>')
    extra = ''
    if id(sl) in AUD:
        sid, dur = AUD.pop(id(sl))
        pars.append(f'<p:par><p:cTn id="{nid()}" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="{"withEffect" if pars else "afterEffect"}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                    f'<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="{nid()}" dur="{dur}" fill="hold"/>{tgt(sid)}</p:cBhvr></p:cmd></p:childTnLst></p:cTn></p:par>')
        extra = (f'<p:audio><p:cMediaNode vol="80000" showWhenStopped="0"><p:cTn id="{nid()}" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn>'
                 f'{tgt(sid)}</p:cMediaNode></p:audio>')
    sl._element.append(etree.fromstring(f'<p:transition {NS} spd="med">{trans}</p:transition>'))
    if not pars: return
    sl._element.append(etree.fromstring(
        f'<p:timing {NS}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
        f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
        f'<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
        f'<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{"".join(pars)}</p:childTnLst></p:cTn></p:par>'
        f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>'
        f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>{extra}'
        f'</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'))
